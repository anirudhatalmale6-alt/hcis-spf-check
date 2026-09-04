#!/usr/bin/env python3
"""
spfcheck - check an HCIS "SPF DED" monthly file before it goes to Finance.

Usage:   python3 spfcheck.py "SPF DED SPF AUGUST 2026.xlsx"

It reads the file and nothing else. It never writes to it, never sends it
anywhere, and never changes a figure. It just prints what it found.

Written for DataBytes Consulting / HCIS, September 2026.
"""

import sys, collections, datetime, calendar

try:
    import openpyxl
except ImportError:
    sys.exit("This needs openpyxl.  Install it with:  pip install openpyxl")

COLS = ['NIN', 'Fullname', 'HomeCareType', 'BasicPay', 'Allowances',
        'Deductions', 'NetPay', 'SPFALLOWABLEPAY', 'SPFEmployee',
        'SPFEmployer', 'Year', 'Month']

EPOCH = datetime.datetime(1899, 12, 30)   # Excel's day zero
SPF_RATE = 0.05
MONEY = 0.01

problems = []
notes = []


def money(v):
    """A figure, however Excel has decided to store it this week.

    A cell formatted as a date still holds a plain number underneath, so a
    date here is not a date - it is an amount wearing the wrong format.
    """
    if v is None or v == '':
        return 0.0
    if isinstance(v, bool):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, datetime.datetime):
        return (v - EPOCH).total_seconds() / 86400.0
    if isinstance(v, datetime.time):
        return (v.hour * 3600 + v.minute * 60 + v.second) / 86400.0
    try:
        return float(str(v).replace(',', '').strip())
    except ValueError:
        return None


def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = [c.value for c in ws[1]][:len(COLS)]
    if [str(h).strip() if h else h for h in header] != COLS:
        problems.append(
            'The columns are not the ones Finance expect.\n'
            '      expected: ' + ', '.join(COLS) + '\n'
            '      found:    ' + ', '.join(str(h) for h in header))
    rows, blanks = [], 0
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if r[0] in (None, ''):
            if any(c not in (None, '') for c in r):
                problems.append('Row %d has no NIN but does have other data '
                                'in it.' % i)
            else:
                blanks += 1
            continue
        rows.append((i, r))
    return ws, rows, blanks


def check_duplicates(rows):
    seen = collections.defaultdict(list)
    for i, r in rows:
        seen[r[0]].append((i, r))
    for nin, group in seen.items():
        if len(group) == 1:
            continue
        same = len(set(tuple(str(x) for x in r) for _, r in group)) == 1
        lines = ', '.join(str(i) for i, _ in group)
        extra_net = sum(money(r[6]) for _, r in group[1:])
        extra_spf = sum(money(r[8]) for _, r in group[1:])
        problems.append(
            '%s appears %d times (rows %s)%s.\n'
            '      If the receiving system adds up rows, that is %.2f of net pay '
            'and %.2f of SPF employee\n'
            '      contribution more than this person should have.'
            % (nin, len(group), lines,
               ', and every copy is identical' if same else '',
               extra_net, extra_spf))


def check_arithmetic(rows):
    bad = 0
    for i, r in rows:
        b, a, d, n = money(r[3]), money(r[4]), money(r[5]), money(r[6])
        if None in (b, a, d, n):
            problems.append('Row %d has a figure that is not a number.' % i)
            continue
        if abs(round(b + a - d - n, 2)) > MONEY:
            bad += 1
            if bad <= 10:
                problems.append(
                    'Row %d does not add up: %.2f + %.2f - %.2f should be %.2f, '
                    'the file says %.2f.' % (i, b, a, d, b + a - d, n))
    if bad > 10:
        problems.append('...and %d more rows that do not add up.' % (bad - 10))
    if not bad:
        notes.append('Every row adds up: basic + allowances - deductions = net.')


def check_spf(rows):
    off = []
    for i, r in rows:
        ap, ee, er = money(r[7]), money(r[8]), money(r[9])
        if abs(ee - er) > MONEY:
            problems.append('Row %d: the employee share (%.2f) and the employer '
                            'share (%.2f) are not the same.' % (i, ee, er))
        if ap and abs(ee - round(ap * SPF_RATE)) > 1.0:
            off.append((i, ap, ee))
    for i, ap, ee in off[:10]:
        problems.append('Row %d: 5%% of %.0f is %.0f, the file says %.0f.'
                        % (i, ap, round(ap * SPF_RATE), ee))
    if len(off) > 10:
        problems.append('...and %d more rows where the 5%% does not come out.'
                        % (len(off) - 10))
    if not off:
        notes.append('The pension contribution is 5% of allowable pay on every '
                     'row, employee and employer alike.')


def check_deduction_days(rows):
    """A deduction that lowers pensionable pay is pay not earned - days absent.

    Nothing in the file says how many days. This works it out backwards, and
    says so when a figure does not land on a whole day, because that is worth
    a second look before the file leaves.
    """
    periods = collections.Counter((r[10], r[11]) for _, r in rows)
    if len(periods) > 1:
        problems.append('This file covers more than one month: '
                        + ', '.join('%s/%s' % p for p in sorted(periods)))
    (year, month), _ = periods.most_common(1)[0]
    try:
        days = calendar.monthrange(int(year), int(month))[1]
    except (TypeError, ValueError):
        notes.append('Could not read the year and month, so days absent were '
                     'not checked.')
        return

    reduced, whole, odd = 0, 0, []
    for i, r in rows:
        gap = money(r[3]) + money(r[4]) - money(r[7])
        if gap <= 1.0:
            continue
        reduced += 1
        rate = money(r[3]) / days
        n = gap / rate
        if abs(n - round(n)) < 0.03 and round(n) >= 1:
            whole += 1
        else:
            odd.append((i, gap, rate))

    notes.append('%d of %d people had pay withheld this month (%.1f%%). '
                 'A day is worth basic pay / %d.'
                 % (reduced, len(rows), 100.0 * reduced / max(1, len(rows)), days))
    if whole:
        notes.append('%d of those %d land exactly on a whole number of days.'
                     % (whole, reduced))
    if odd:
        problems.append(
            '%d rows withheld an amount that is not a whole number of days at '
            'basic / %d.\n'
            '      That is not necessarily wrong - a flat deduction, say a '
            'salary advance being repaid,\n'
            '      would look exactly like this. But the file gives no reason '
            'for any deduction, so\n'
            '      nobody reading it can tell the two apart. Worth a look:'
            % (len(odd), days))
    for i, gap, rate in odd[:10]:
        problems.append('    row %d: %.2f withheld = %.2f days at %.2f a day'
                        % (i, gap, gap / rate, rate))
    if len(odd) > 10:
        problems.append('    ...and %d more.' % (len(odd) - 10))


def check_formats(ws, rows):
    """A money column wearing a date format.

    The number underneath is right - that is why the rows still add up - but
    anyone opening the file sees a date from 1903 where an allowance should be,
    and an importer that reads the displayed value rather than the stored one
    gets a date too.
    """
    for col, name in ((5, 'Allowances'), (4, 'BasicPay'), (6, 'Deductions'),
                      (7, 'NetPay')):
        fmts = collections.Counter(ws.cell(row=i, column=col).number_format
                                   for i, _ in rows[:500])
        for fmt, n in fmts.items():
            if any(ch in str(fmt) for ch in ('y', 'd/m', 'h:m')):
                problems.append(
                    'The %s column is formatted as a date or a time (%s). '
                    'It holds money.\n'
                    '      Set that column to Number with 2 decimals before '
                    'sending the file.' % (name, fmt))
                break


def check_blank_tail(ws, rows, blanks):
    if blanks > 100:
        problems.append(
            'There are %d empty rows underneath the data. The sheet claims to '
            'be %d rows long\n'
            '      when only %d of them are people. Delete the empty rows '
            '(select them, right-click, Delete)\n'
            '      or a system reading to the end of the sheet will try to '
            'import %d nothings.'
            % (blanks, ws.max_row, len(rows), blanks))


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__.strip())
    path = sys.argv[1]
    ws, rows, blanks = load(path)
    if not rows:
        sys.exit('No rows with a NIN in them - is this the right file?')

    check_duplicates(rows)
    check_arithmetic(rows)
    check_spf(rows)
    check_deduction_days(rows)
    check_formats(ws, rows)
    check_blank_tail(ws, rows, blanks)

    total = lambda c: sum(money(r[c]) for _, r in rows)
    print('\n%s' % path)
    print('%d people, %.2f basic, %.2f allowances, %.2f deductions, %.2f net.'
          % (len(rows), total(3), total(4), total(5), total(6)))
    print('SPF: %.2f from the employee, %.2f from HC.' % (total(8), total(9)))

    if notes:
        print('\nChecked and fine:')
        for n in notes:
            print('  - %s' % n)

    if problems:
        print('\nWorth looking at before this goes to Finance (%d):' % len(problems))
        for p in problems:
            print('  ! %s' % p)
        print('')
        return 1
    print('\nNothing to flag.\n')
    return 0


if __name__ == '__main__':
    sys.exit(main())
